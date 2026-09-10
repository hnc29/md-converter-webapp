from pathlib import Path
from typing import List, Optional
import openpyxl
import pandas as pd

class TableProcessor:
    """
    Handles robust tabular data conversion for Word tables and Excel spreadsheets,
    ensuring searchable, clean Markdown tables.
    """

    @staticmethod
    def format_markdown_table(rows: List[List[str]]) -> str:
        """Converts a 2D matrix of strings into a standard Markdown table."""
        if not rows:
            return ""

        # Normalize cells: replace newlines with space, strip whitespace, escape pipes
        clean_rows = []
        for row in rows:
            clean_row = [str(cell if cell is not None else '').replace('\n', ' ').replace('\r', ' ').replace('|', '\\|').strip() for cell in row]
            clean_rows.append(clean_row)

        if not clean_rows:
            return ""

        # Normalize column count to max columns
        max_cols = max(len(r) for r in clean_rows)
        if max_cols == 0:
            return ""

        padded_rows = []
        for r in clean_rows:
            padded = r + [''] * (max_cols - len(r))
            padded_rows.append(padded)

        header = padded_rows[0]
        # If header is all empty, generate Col 1, Col 2...
        if all(c == '' for c in header):
            header = [f"Cột {i+1}" for i in range(max_cols)]
            data_rows = padded_rows
        else:
            data_rows = padded_rows[1:]

        md_lines = []
        md_lines.append(f"| {' | '.join(header)} |")
        md_lines.append(f"| {' | '.join(['---'] * max_cols)} |")
        for r in data_rows:
            md_lines.append(f"| {' | '.join(r)} |")

        return "\n".join(md_lines)

    def convert_excel_to_markdown(self, excel_path: Path, include_formula: bool = False) -> str:
        """
        Converts an Excel workbook (multi-sheet) into a clean, hierarchical Markdown document.
        Each sheet becomes a section: ## Sheet: <SheetName>
        """
        filename = excel_path.name
        md_parts = [f"# File: {filename}\n"]

        try:
            # Try openpyxl first for .xlsx
            wb = openpyxl.load_workbook(excel_path, data_only=not include_formula)
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                raw_rows = list(ws.iter_rows(values_only=True))
                if not raw_rows:
                    continue

                # Filter completely empty rows
                valid_rows = [[str(cell if cell is not None else '').strip() for cell in r] for r in raw_rows if any(cell is not None and str(cell).strip() != '' for cell in r)]
                if not valid_rows:
                    continue

                md_parts.append(f"## Sheet: {sheet_name}\n")
                tbl_md = self.format_markdown_table(valid_rows)
                md_parts.append(tbl_md)
                md_parts.append("\n")

        except Exception:
            # Fallback to pandas with xlrd/openpyxl
            excel_file = pd.ExcelFile(excel_path)
            for sheet_name in excel_file.sheet_names:
                df = pd.read_excel(excel_file, sheet_name=sheet_name, header=None)
                if df.empty:
                    continue
                valid_rows = df.fillna('').astype(str).values.tolist()
                md_parts.append(f"## Sheet: {sheet_name}\n")
                tbl_md = self.format_markdown_table(valid_rows)
                md_parts.append(tbl_md)
                md_parts.append("\n")

        return "\n".join(md_parts).strip()

table_processor = TableProcessor()
