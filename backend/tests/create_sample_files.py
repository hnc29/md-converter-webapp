import io
import subprocess
from pathlib import Path
from docx import Document
import openpyxl
import pymupdf
from PIL import Image, ImageDraw

TESTS_DIR = Path(__file__).resolve().parent
FIXTURES_DIR = TESTS_DIR / "fixtures"
FIXTURES_DIR.mkdir(parents=True, exist_ok=True)

def create_ims_hld_docx():
    doc_path = FIXTURES_DIR / "03.IMS_HLD.docx"
    doc = Document()
    doc.core_properties.title = "Tài liệu thiết kế tổng thể hệ thống IMS"
    doc.core_properties.author = "VNPT"

    doc.add_heading("TÀI LIỆU THIẾT KẾ TỔNG THỂ HỆ THỐNG IMS", level=1)
    doc.add_paragraph("Tập đoàn Bưu chính Viễn thông Việt Nam (VNPT)")
    doc.add_paragraph("Phiên bản: 1.0 | Trạng thái: Đã ban hành | Ngày: 26/08/2026")

    doc.add_heading("1. Tổng quan và Phạm vi", level=2)
    doc.add_paragraph("Hệ thống Quản lý Đầu tư Xây dựng (IMS) là hệ thống cốt lõi quản lý toàn trình dự án đầu tư.")

    doc.add_heading("2. Kiến trúc hệ thống", level=2)
    doc.add_heading("2.1. Kiến trúc ứng dụng", level=3)
    doc.add_paragraph("Kiến trúc ứng dụng phân tầng Microservices kết nối qua API Gateway.")
    doc.add_heading("2.2. Kiến trúc tích hợp", level=3)
    doc.add_paragraph("Hệ thống IMS tích hợp với SAP ERP, CRM và PMIS thông qua Enterprise Service Bus (ESB).")

    doc.add_heading("3. Các phân hệ chức năng", level=2)
    doc.add_heading("3.1. Quản lý chủ trương đầu tư", level=3)
    doc.add_paragraph("Tiếp nhận và thẩm định chủ trương đầu tư các dự án hạ tầng mạng viễn thông.")
    doc.add_heading("3.2. Quản lý dự án", level=3)
    doc.add_paragraph("Theo dõi tiến độ thực hiện dự án, quản lý hợp đồng và nghiệm thu kỹ thuật.")
    doc.add_heading("3.3. Quản lý kế hoạch vốn", level=3)
    doc.add_paragraph("IMS có chức năng quản lý kế hoạch vốn đầu tư trung hạn và hàng năm, theo dõi giải ngân thanh toán vốn đầu tư.")

    # Table
    table = doc.add_table(rows=1, cols=4)
    hdr = table.rows[0].cells
    hdr[0].text = "STT"
    hdr[1].text = "Chức năng"
    hdr[2].text = "Mô tả nghiệp vụ"
    hdr[3].text = "Đơn vị thực hiện"

    items = [
        ("1", "Quản lý dự án", "Quản lý thông tin và tiến độ dự án", "KHĐT"),
        ("2", "Kế hoạch vốn", "Theo dõi kế hoạch vốn và thanh toán", "KTTC"),
        ("3", "Nghiệm thu quyết toán", "Hồ sơ nghiệm thu bàn giao tài sản", "Ban QLDA")
    ]
    for stt, func, desc, unit in items:
        row = table.add_row().cells
        row[0].text = stt
        row[1].text = func
        row[2].text = desc
        row[3].text = unit

    doc.save(doc_path)
    return doc_path

def create_ims_lld_docx():
    doc_path = FIXTURES_DIR / "04.IMS_LLD.docx"
    doc = Document()
    doc.core_properties.title = "Tài liệu thiết kế chi tiết hệ thống IMS"
    doc.add_heading("TÀI LIỆU THIẾT KẾ CHI TIẾT HỆ THỐNG IMS", level=1)
    doc.add_paragraph("Phiên bản: 1.0 | Mã tài liệu: IMS-LLD")

    doc.add_heading("1. Thiết kế cơ sở dữ liệu", level=2)
    doc.add_paragraph("Thiết kế lược đồ quan hệ PostgreSQL và chỉ mục hiệu năng cao.")

    doc.add_heading("2. Thiết kế API", level=2)
    doc.add_paragraph("Đặc tả chi tiết các RESTful API và mô hình xác thực JWT.")

    doc.save(doc_path)
    return doc_path

def create_ims_excel_test_report():
    xlsx_path = FIXTURES_DIR / "IMS_TEST_REPORT_PhiChucNang.xlsx"
    wb = openpyxl.Workbook()

    ws1 = wb.active
    ws1.title = "TongHop"
    ws1.append(["Ma Kiem Thu", "Hang Muc", "Ket Qua", "Ghi Chu"])
    ws1.append(["TC01", "Kiem thu tai tong the", "Dat", "1000 concurrent users"])
    ws1.append(["TC02", "Kiem thu an toan thong tin", "Dat", "Vượt qua đánh giá SOC"])

    ws2 = wb.create_sheet(title="Performance")
    ws2.append(["Chi Tieu", "Nguong Yeu Cau", "Thuc Te", "Danh Gia"])
    ws2.append(["Thoi gian dap ung", "< 200ms", "145ms", "Dat"])
    ws2.append(["Throughput", "> 1000 TPS", "1250 TPS", "Dat"])

    wb.save(xlsx_path)
    return xlsx_path

def create_sample_legacy_doc():
    doc_path = FIXTURES_DIR / "sample_legacy.doc"
    txt_path = FIXTURES_DIR / "temp_sample.txt"
    txt_path.write_text("Tai lieu Word cu dinh dang DOC thu nghiem chuyen doi Markdown.", encoding="utf-8")
    subprocess.run(["textutil", "-convert", "doc", "-output", str(doc_path), str(txt_path)], check=True)
    if txt_path.exists():
        txt_path.unlink()
    return doc_path

def create_sample_legacy_xls():
    xls_path = FIXTURES_DIR / "sample_legacy.xls"
    import xlwt
    wb = xlwt.Workbook()
    ws1 = wb.add_sheet("BangGia")
    ws1.write(0, 0, "San Pham")
    ws1.write(0, 1, "Gia Ban")
    ws1.write(1, 0, "Goi Cuoc 4G")
    ws1.write(1, 1, "120000")
    wb.save(str(xls_path))
    return xls_path

def create_sample_digital_pdf():
    pdf_path = FIXTURES_DIR / "sample_digital.pdf"
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    text = (
        "CONG HOA XA HOI CHU NGHIA VIET NAM\n"
        "Doc lap - Tu do - Hanh phuc\n\n"
        "BAO CAO KET QUA TRIEN KHAI HE THONG KNOWLEDGE BASE\n\n"
        "Tai lieu nay chua text layer so ro rang va day du.\n"
        "He thong tu dong trich xuat va chen marker ranh gioi trang source_page."
    )
    page.insert_text((50, 72), text, fontsize=11)
    doc.save(str(pdf_path))
    doc.close()
    return pdf_path

def create_sample_scanned_pdf():
    pdf_path = FIXTURES_DIR / "sample_scanned.pdf"
    img = Image.new("RGB", (800, 600), color=(255, 255, 255))
    draw = ImageDraw.Draw(img)
    draw.text((50, 50), "TAI LIEU SCAN ANH (OCR LOCAL TEST)", fill=(0, 0, 0))
    draw.text((50, 100), "Kiem tra nhan dien OCR va gan nhan [Extracted from image]", fill=(0, 0, 0))

    img_bytes = io.BytesIO()
    img.save(img_bytes, format="PNG")

    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    rect = pymupdf.Rect(50, 50, 545, 420)
    page.insert_image(rect, stream=img_bytes.getvalue())
    doc.save(str(pdf_path))
    doc.close()
    return pdf_path

def setup_all_fixtures():
    create_ims_hld_docx()
    create_ims_lld_docx()
    create_ims_excel_test_report()
    create_sample_legacy_doc()
    create_sample_legacy_xls()
    create_sample_digital_pdf()
    create_sample_scanned_pdf()

# Backward-compatibility aliases
create_sample_docx = create_ims_hld_docx
create_sample_xlsx = create_ims_excel_test_report
create_sample_doc = create_sample_legacy_doc
create_sample_xls = create_sample_legacy_xls
create_sample_mixed_pdf = create_sample_digital_pdf

if __name__ == "__main__":
    setup_all_fixtures()
    print("All fixtures created successfully!")
